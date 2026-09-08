"""
Iteration 56 backend tests:
- PATCH /api/sites/{id} accepting company_id + city_id (master only)
- PATCH /api/master/corporate-clients/{id} accepting city_id
- GET /api/admin/sales-report/filters cascading data
- GET /api/admin/sales-report filtered totals (seeded DEMO June 2026)
- Excel export with 4 sheets
- RBAC: non-master gets 403 on site company_id PATCH and admin sales report
- Cleanup test-created rows
"""
import os
import io
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@cravitoo.com"
ADMIN_PASS = "admin123"


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_headers():
    return {"Authorization": f"Bearer {_login(ADMIN_EMAIL, ADMIN_PASS)}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def created_ids():
    return {"client_id": None, "city_id": None, "site_id": None}


# ---------- SITE LINK REPAIR ----------

def test_create_client_city_site_and_link(admin_headers, created_ids):
    import uuid
    suffix = uuid.uuid4().hex[:6]

    # 1. Create corporate client
    client_payload = {
        "name": f"TEST_Client_{suffix}",
        "address": "1 Test Rd",
        "contact_email": f"test_client_{suffix}@example.com",
        "contact_phone": "9999999999",
    }
    r = requests.post(f"{API}/master/corporate-clients", headers=admin_headers, json=client_payload, timeout=30)
    assert r.status_code in (200, 201), f"create client: {r.status_code} {r.text}"
    client = r.json()
    client_id = client.get("id") or client.get("_id")
    assert client_id, f"no id in {client}"
    created_ids["client_id"] = client_id

    # 2. Create city
    city_payload = {"name": f"TEST_City_{suffix}", "state": "TS"}
    r = requests.post(f"{API}/cities", headers=admin_headers, json=city_payload, timeout=30)
    assert r.status_code in (200, 201), f"create city: {r.status_code} {r.text}"
    city = r.json()
    city_id = city.get("id") or city.get("_id")
    assert city_id
    created_ids["city_id"] = city_id

    # 3. Create site WITHOUT company_id / city_id
    site_payload = {
        "name": f"TEST_Site_{suffix}",
        "address": "1 Site Ln",
        "city": "Testville",
        "contact_email": f"test_site_{suffix}@example.com",
        "contact_phone": "8888888888",
    }
    r = requests.post(f"{API}/sites", headers=admin_headers, json=site_payload, timeout=30)
    assert r.status_code in (200, 201), f"create site: {r.status_code} {r.text}"
    site = r.json()
    site_id = site.get("id") or site.get("_id")
    assert site_id
    created_ids["site_id"] = site_id
    assert not site.get("company_id"), f"expected no company_id on fresh site, got {site.get('company_id')}"

    # 4. PATCH site with company_id + city_id
    patch_payload = {"company_id": client_id, "city_id": city_id}
    r = requests.patch(f"{API}/sites/{site_id}", headers=admin_headers, json=patch_payload, timeout=30)
    assert r.status_code == 200, f"patch site: {r.status_code} {r.text}"

    # 5. GET site verify persistence
    r = requests.get(f"{API}/sites/{site_id}", headers=admin_headers, timeout=30)
    assert r.status_code == 200, r.text
    got = r.json()
    assert got.get("company_id") == client_id, f"company_id not persisted: {got.get('company_id')}"
    assert got.get("city_id") == city_id, f"city_id not persisted: {got.get('city_id')}"


def test_client_city_patch(admin_headers, created_ids):
    client_id = created_ids["client_id"]
    city_id = created_ids["city_id"]
    assert client_id and city_id
    r = requests.patch(
        f"{API}/master/corporate-clients/{client_id}",
        headers=admin_headers,
        json={"city_id": city_id},
        timeout=30,
    )
    assert r.status_code == 200, f"patch client city_id: {r.status_code} {r.text}"

    # verify in list
    r = requests.get(f"{API}/master/corporate-clients", headers=admin_headers, timeout=30)
    assert r.status_code == 200
    items = r.json() if isinstance(r.json(), list) else r.json().get("clients", r.json().get("data", []))
    found = next((c for c in items if (c.get("id") or c.get("_id")) == client_id), None)
    assert found, f"client {client_id} not in list"
    assert found.get("city_id") == city_id, f"city_id not persisted on client: {found.get('city_id')}"


# ---------- CASCADE FILTERS ----------

def test_filters_endpoint(admin_headers):
    r = requests.get(f"{API}/admin/sales-report/filters", headers=admin_headers, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    for key in ("clients", "cities", "sites", "vendors"):
        assert key in data, f"missing key {key} in filters response"

    # every site should have company_id + city_id keys (may be None but must exist)
    for s in data["sites"]:
        assert "company_id" in s and "city_id" in s, f"site missing keys: {s}"
    for v in data["vendors"]:
        assert "site_ids" in v, f"vendor missing site_ids: {v}"

    # DEMO seeded rows must exist
    client_names = [c.get("name") for c in data["clients"]]
    city_names = [c.get("name") for c in data["cities"]]
    assert any("DEMO Acme Corp" in n for n in client_names if n), f"DEMO Acme missing: {client_names}"
    assert any("DEMO Globex" in n for n in client_names if n), f"DEMO Globex missing: {client_names}"
    assert any("DEMO Bengaluru" in n for n in city_names if n), f"DEMO Bengaluru missing: {city_names}"
    assert any("DEMO Mumbai" in n for n in city_names if n), f"DEMO Mumbai missing: {city_names}"


# ---------- FILTERED REPORT ----------

def _find_id_by_name(items, name_substr):
    for it in items:
        if it.get("name") and name_substr in it["name"]:
            return it.get("id") or it.get("_id")
    return None


@pytest.fixture(scope="module")
def demo_ids(admin_headers):
    r = requests.get(f"{API}/admin/sales-report/filters", headers=admin_headers, timeout=30)
    assert r.status_code == 200
    d = r.json()
    return {
        "acme": _find_id_by_name(d["clients"], "DEMO Acme Corp"),
        "globex": _find_id_by_name(d["clients"], "DEMO Globex"),
        "bengaluru": _find_id_by_name(d["cities"], "DEMO Bengaluru"),
        "mumbai": _find_id_by_name(d["cities"], "DEMO Mumbai"),
    }


def _fetch_report(headers, **params):
    r = requests.get(f"{API}/admin/sales-report", headers=headers, params=params, timeout=60)
    assert r.status_code == 200, f"{params} -> {r.status_code} {r.text}"
    return r.json()


def test_report_grand_total(admin_headers):
    d = _fetch_report(admin_headers, month="2026-06")
    assert float(d.get("grand_total", 0)) == pytest.approx(3975.00, abs=0.01), f"grand_total={d.get('grand_total')}"
    assert d.get("total_orders") == 14 or d.get("orders_count") == 14 or len(d.get("orders", [])) == 14, (
        f"expected 14 orders, got {d.get('total_orders') or d.get('orders_count') or len(d.get('orders', []))}"
    )
    for key in ("client_summary", "city_summary", "site_summary", "vendor_summary"):
        assert key in d, f"missing {key} in report"


def test_report_client_acme(admin_headers, demo_ids):
    assert demo_ids["acme"], "Acme id not found"
    d = _fetch_report(admin_headers, month="2026-06", client_ids=demo_ids["acme"])
    assert float(d["grand_total"]) == pytest.approx(2390.00, abs=0.01), f"acme total={d['grand_total']}"
    n = d.get("total_orders") or d.get("orders_count") or len(d.get("orders", []))
    assert n == 8, f"acme orders={n}"


def test_report_city_bengaluru(admin_headers, demo_ids):
    assert demo_ids["bengaluru"]
    d = _fetch_report(admin_headers, month="2026-06", city_ids=demo_ids["bengaluru"])
    assert float(d["grand_total"]) == pytest.approx(2505.00, abs=0.01), f"blr total={d['grand_total']}"
    n = d.get("total_orders") or d.get("orders_count") or len(d.get("orders", []))
    assert n == 9


def test_report_acme_mumbai(admin_headers, demo_ids):
    d = _fetch_report(admin_headers, month="2026-06", client_ids=demo_ids["acme"], city_ids=demo_ids["mumbai"])
    assert float(d["grand_total"]) == pytest.approx(1170.00, abs=0.01), f"acme+mumbai={d['grand_total']}"
    n = d.get("total_orders") or d.get("orders_count") or len(d.get("orders", []))
    assert n == 3


def test_report_multi_client(admin_headers, demo_ids):
    ids = f"{demo_ids['acme']},{demo_ids['globex']}"
    d = _fetch_report(admin_headers, month="2026-06", client_ids=ids)
    assert float(d["grand_total"]) == pytest.approx(3675.00, abs=0.01), f"multi={d['grand_total']}"
    n = d.get("total_orders") or d.get("orders_count") or len(d.get("orders", []))
    assert n == 12


# ---------- EXCEL ----------

def test_excel_export_four_sheets(admin_headers, demo_ids):
    r = requests.get(
        f"{API}/admin/sales-report",
        headers=admin_headers,
        params={"month": "2026-06", "client_ids": demo_ids["acme"], "format": "xlsx"},
        timeout=60,
    )
    assert r.status_code == 200, r.text
    ct = r.headers.get("content-type", "")
    assert "spreadsheet" in ct or "xlsx" in ct or "octet-stream" in ct, f"content-type: {ct}"
    try:
        from openpyxl import load_workbook
    except ImportError:
        pytest.skip("openpyxl not installed")
    wb = load_workbook(io.BytesIO(r.content))
    names = set(wb.sheetnames)
    expected = {"Orders", "City Totals", "Site Totals", "Vendor Totals"}
    assert expected.issubset(names), f"missing sheets. got {names}"


# ---------- RBAC ----------

def test_rbac_non_master_forbidden(created_ids, admin_headers):
    # Provision a fresh employee on an allowed domain (linktest.com) via self-register
    import uuid
    email = f"test_rbac_{uuid.uuid4().hex[:6]}@linktest.com"
    password = "RbacTest#123"
    rr = requests.post(
        f"{API}/auth/register",
        json={"email": email, "password": password, "name": "RBAC Test", "role": "employee"},
        timeout=30,
    )
    assert rr.status_code in (200, 201), f"register employee failed: {rr.status_code} {rr.text}"
    tok = _login(email, password)
    h = {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}
    # Track user id for cleanup
    me = requests.get(f"{API}/auth/me", headers=h, timeout=30)
    if me.status_code == 200:
        created_ids["rbac_user_id"] = me.json().get("id") or me.json().get("_id") or me.json().get("user", {}).get("id")

    r = requests.patch(f"{API}/sites/{created_ids['site_id']}", headers=h, json={"company_id": "x"}, timeout=30)
    assert r.status_code == 403, f"employee patch site company_id expected 403, got {r.status_code} {r.text}"

    r = requests.get(f"{API}/admin/sales-report", headers=h, params={"month": "2026-06"}, timeout=30)
    assert r.status_code == 403, f"employee sales-report expected 403, got {r.status_code}"

    r = requests.get(f"{API}/admin/sales-report/filters", headers=h, timeout=30)
    assert r.status_code == 403, f"employee filters expected 403, got {r.status_code}"


# ---------- CLEANUP ----------

def test_zz_cleanup(admin_headers, created_ids):
    # delete client with cascade
    cid = created_ids.get("client_id")
    site_id = created_ids.get("site_id")
    city_id = created_ids.get("city_id")

    if site_id:
        r = requests.delete(f"{API}/sites/{site_id}", headers=admin_headers, timeout=30)
        print(f"delete site -> {r.status_code}")
    if cid:
        r = requests.delete(f"{API}/master/corporate-clients/{cid}?cascade=true", headers=admin_headers, timeout=30)
        print(f"delete client -> {r.status_code}")
    if city_id:
        r = requests.delete(f"{API}/cities/{city_id}", headers=admin_headers, timeout=30)
        print(f"delete city -> {r.status_code}")
    rbac_uid = created_ids.get("rbac_user_id")
    if rbac_uid:
        r = requests.post(f"{API}/admin/users/{rbac_uid}/deactivate", headers=admin_headers, timeout=30)
        print(f"deactivate rbac user -> {r.status_code}")
