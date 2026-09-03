"""Iter 35: Backend tests for 4 menu-management additions:
   1) Excel Template Download (GET /api/admin/menu-excel-template)
   2) Upload Preview / Diff (POST /api/sites/{id}/menu/preview)
   3) Menu Version History + Restore
   4) Vendor Bulk Upload → Admin Approval
"""
import io
import os
import time
import pytest
import requests
from openpyxl import Workbook, load_workbook


def _load_env_url():
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if v:
        return v.rstrip("/")
    with open("/app/frontend/.env") as f:
        for ln in f:
            if ln.startswith("REACT_APP_BACKEND_URL="):
                return ln.strip().split("=", 1)[1].rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL not found")


BASE_URL = _load_env_url()
API = f"{BASE_URL}/api"

SITE_A = "6a96d37822127a5e04873e36"
VENDOR_1 = "6a96d37822127a5e04873e3c"

CREDS = {
    "master": ("admin@cravitoo.com", "admin123"),
    "vendor1": ("audit_vendor1@corpa.com", "Audit#1234"),
    "emp_a": ("audit_emp_a@corpa.com", "Audit#1234"),
}


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed {email}: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def tokens():
    return {k: _login(*v) for k, v in CREDS.items()}


def _h(t):
    return {"Authorization": f"Bearer {t}"}


def _live_menu(tok):
    r = requests.get(f"{API}/menu/{VENDOR_1}", headers=_h(tok), timeout=30)
    assert r.status_code == 200, r.text
    return r.json()


def _make_xlsx(rows):
    """rows: list of dicts w/ name, description, category, price, is_vegetarian, meal_periods"""
    wb = Workbook()
    ws = wb.active
    headers = ["name", "description", "category", "price", "is_vegetarian", "image_url", "meal_periods"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------- Feature 1: Excel template download ----------
class TestTemplateDownload:
    def test_master_download(self, tokens):
        r = requests.get(f"{API}/admin/menu-excel-template", headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert len(r.content) > 3000, f"template too small: {len(r.content)}"
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for col in ("name", "description", "category", "price", "is_vegetarian"):
            assert col in headers, f"missing column {col}, got {headers}"

    def test_vendor_can_download(self, tokens):
        r = requests.get(f"{API}/admin/menu-excel-template", headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 200, r.text
        assert len(r.content) > 3000

    def test_unauth_forbidden(self):
        r = requests.get(f"{API}/admin/menu-excel-template", timeout=30)
        assert r.status_code in (401, 403)


# ---------- Feature 2: Upload preview (diff) ----------
class TestUploadPreview:
    def test_preview_no_mutation(self, tokens):
        # Fetch current live count via master API
        before = _live_menu(tokens["master"])
        before_names = {i["name"] for i in before}
        rows = [
            {"name": "PREVIEW_ItemA", "description": "d", "category": "Snacks", "price": 50, "is_vegetarian": True},
            {"name": "PREVIEW_ItemB", "description": "d", "category": "Snacks", "price": 60, "is_vegetarian": True},
        ]
        # include one existing name (if any) to trigger 'updated'
        if before_names:
            rows.append({"name": next(iter(before_names)), "description": "d", "category": "Snacks", "price": 99, "is_vegetarian": True})
        xlsx = _make_xlsx(rows)
        files = {"file": ("preview.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(
            f"{API}/sites/{SITE_A}/menu/preview",
            params={"vendor_id": VENDOR_1},
            headers=_h(tokens["master"]), files=files, timeout=30,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        for k in ("added", "updated", "removed", "total_in_file", "current_count", "errors"):
            assert k in data, f"missing {k}"
        assert "PREVIEW_ItemA" in data["added"]
        assert "PREVIEW_ItemB" in data["added"]
        assert data["total_in_file"] == len(rows)
        # verify live menu unchanged
        after = _live_menu(tokens["master"])
        after_names = {i["name"] for i in after}
        assert "PREVIEW_ItemA" not in after_names, "preview must NOT change live menu"
        assert "PREVIEW_ItemB" not in after_names

    def test_preview_vendor_forbidden(self, tokens):
        # vendors don't have site-admin access → 403
        xlsx = _make_xlsx([{"name": "X", "category": "Snacks", "price": 10, "is_vegetarian": True}])
        files = {"file": ("x.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(
            f"{API}/sites/{SITE_A}/menu/preview",
            params={"vendor_id": VENDOR_1},
            headers=_h(tokens["vendor1"]), files=files, timeout=30,
        )
        assert r.status_code == 403, r.text

    def test_preview_employee_forbidden(self, tokens):
        xlsx = _make_xlsx([{"name": "X", "category": "Snacks", "price": 10, "is_vegetarian": True}])
        files = {"file": ("x.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(
            f"{API}/sites/{SITE_A}/menu/preview",
            params={"vendor_id": VENDOR_1},
            headers=_h(tokens["emp_a"]), files=files, timeout=30,
        )
        assert r.status_code == 403


# ---------- Feature 4: Vendor Bulk Upload → Admin Approval (do this before Feature 3 so a version snapshot is created) ----------
class TestVendorBulkApprovalFlow:
    @pytest.fixture(scope="class")
    def submission_id(self, tokens):
        rows = [
            {"name": "APPROVAL_Idli", "description": "steamed", "category": "Breakfast", "price": 40, "is_vegetarian": True},
            {"name": "APPROVAL_Vada", "description": "fried", "category": "Breakfast", "price": 45, "is_vegetarian": True},
            {"name": "APPROVAL_Dosa", "description": "crisp", "category": "Breakfast", "price": 60, "is_vegetarian": True},
        ]
        xlsx = _make_xlsx(rows)
        files = {"file": ("submit.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        before = _live_menu(tokens["master"])
        before_names = {i["name"] for i in before}
        r = requests.post(
            f"{API}/vendor/menu-uploads",
            params={"site_id": SITE_A},
            headers=_h(tokens["vendor1"]), files=files, timeout=30,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["status"] == "pending"
        assert data["item_count"] == 3
        # live menu untouched by submit
        after = _live_menu(tokens["master"])
        after_names = {i["name"] for i in after}
        assert "APPROVAL_Idli" not in after_names, "vendor submit must NOT change live menu"
        assert after_names == before_names or after_names.issubset(before_names.union(set())), "live menu changed on submit"
        return data["id"]

    def test_vendor_sees_submission_pending(self, tokens, submission_id):
        r = requests.get(f"{API}/vendor/menu-uploads", headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 200
        rows = r.json()
        row = next((x for x in rows if x["id"] == submission_id), None)
        assert row is not None, f"submission not in vendor list: {rows}"
        assert row["status"] == "pending"

    def test_admin_sees_pending(self, tokens, submission_id):
        r = requests.get(f"{API}/admin/menu-uploads", params={"status": "pending"},
                         headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 200, r.text
        rows = r.json()
        row = next((x for x in rows if x["id"] == submission_id), None)
        assert row is not None, f"submission missing in admin list"
        assert row["status"] == "pending"
        assert row["item_count"] == 3

    def test_vendor_cannot_approve(self, tokens, submission_id):
        r = requests.post(f"{API}/admin/menu-uploads/{submission_id}/approve",
                          headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 403

    def test_vendor_cannot_list_admin(self, tokens):
        r = requests.get(f"{API}/admin/menu-uploads", headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 403

    def test_employee_cannot_approve(self, tokens, submission_id):
        r = requests.post(f"{API}/admin/menu-uploads/{submission_id}/approve",
                          headers=_h(tokens["emp_a"]), timeout=30)
        assert r.status_code == 403

    def test_vendor_cannot_upload_to_unmapped_site(self, tokens):
        xlsx = _make_xlsx([{"name": "X", "category": "Snacks", "price": 10, "is_vegetarian": True}])
        files = {"file": ("x.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/vendor/menu-uploads",
                          params={"site_id": "000000000000000000000000"},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        assert r.status_code == 403, r.text

    def test_admin_approve_applies_to_live(self, tokens, submission_id):
        r = requests.post(f"{API}/admin/menu-uploads/{submission_id}/approve",
                          headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["status"] == "approved"
        assert data["applied"] == 3
        # Live menu now contains our 3 items
        live = _live_menu(tokens["master"])
        names = {i["name"] for i in live}
        for n in ("APPROVAL_Idli", "APPROVAL_Vada", "APPROVAL_Dosa"):
            assert n in names, f"{n} missing after approve. live={names}"

    def test_approve_twice_fails(self, tokens, submission_id):
        r = requests.post(f"{API}/admin/menu-uploads/{submission_id}/approve",
                          headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 400  # already approved

    def test_reject_flow_no_mutation(self, tokens):
        # Submit new pending request, then reject and verify live menu unchanged
        rows = [{"name": "REJECT_Item1", "category": "Snacks", "price": 20, "is_vegetarian": True}]
        xlsx = _make_xlsx(rows)
        files = {"file": ("rej.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/vendor/menu-uploads", params={"site_id": SITE_A},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        assert r.status_code == 200
        rid = r.json()["id"]

        before = {i["name"] for i in _live_menu(tokens["master"])}
        rr = requests.post(f"{API}/admin/menu-uploads/{rid}/reject",
                           headers=_h(tokens["master"]),
                           json={"note": "test rejection"}, timeout=30)
        assert rr.status_code == 200, rr.text
        assert rr.json()["status"] == "rejected"
        after = {i["name"] for i in _live_menu(tokens["master"])}
        assert "REJECT_Item1" not in after
        assert before == after, "reject must not modify live menu"


# ---------- Feature 3: Menu versions + restore (after approval above so we have a snapshot) ----------
class TestMenuVersions:
    def test_list_versions_master(self, tokens):
        r = requests.get(f"{API}/sites/{SITE_A}/menu/versions",
                         params={"vendor_id": VENDOR_1},
                         headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert isinstance(data, list)
        assert len(data) >= 1, "at least one snapshot expected from previous approve"
        assert len(data) <= 5, f"should keep only 5 max, got {len(data)}"
        for v in data:
            for k in ("id", "action", "item_count", "created_at"):
                assert k in v

    def test_list_versions_vendor_forbidden(self, tokens):
        r = requests.get(f"{API}/sites/{SITE_A}/menu/versions",
                         params={"vendor_id": VENDOR_1},
                         headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 403

    def test_restore_snapshot(self, tokens):
        # Get current live snapshot names
        live_before = {i["name"] for i in _live_menu(tokens["master"])}

        # Get available versions
        vers = requests.get(f"{API}/sites/{SITE_A}/menu/versions",
                            params={"vendor_id": VENDOR_1},
                            headers=_h(tokens["master"]), timeout=30).json()
        assert vers, "no versions available for restore test"
        vid = vers[0]["id"]  # most recent snapshot

        # Change the live menu to something different via approval flow.
        rows = [{"name": "CHANGE_MarkerItem", "category": "Snacks", "price": 30, "is_vegetarian": True}]
        xlsx = _make_xlsx(rows)
        files = {"file": ("chg.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        s = requests.post(f"{API}/vendor/menu-uploads", params={"site_id": SITE_A},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        rid = s.json()["id"]
        requests.post(f"{API}/admin/menu-uploads/{rid}/approve",
                      headers=_h(tokens["master"]), timeout=30)
        # Now live menu should NOT contain APPROVAL_Idli (replaced)
        mid = {i["name"] for i in _live_menu(tokens["master"])}
        # (site-scoped replace — only items on SiteA changed; menu API is vendor-wide across sites)
        # We just require CHANGE_MarkerItem exists somewhere now
        assert "CHANGE_MarkerItem" in mid

        # Restore vid
        rr = requests.post(f"{API}/sites/{SITE_A}/menu/versions/{vid}/restore",
                           params={"vendor_id": VENDOR_1},
                           headers=_h(tokens["master"]), timeout=30)
        assert rr.status_code == 200, rr.text
        data = rr.json()
        assert "restored" in data
        # Live menu should no longer contain CHANGE_MarkerItem (SiteA was overwritten)
        final = {i["name"] for i in _live_menu(tokens["master"])}
        assert "CHANGE_MarkerItem" not in final, f"restore did not overwrite. final={final}"

    def test_restore_vendor_forbidden(self, tokens):
        vers = requests.get(f"{API}/sites/{SITE_A}/menu/versions",
                            params={"vendor_id": VENDOR_1},
                            headers=_h(tokens["master"]), timeout=30).json()
        if not vers:
            pytest.skip("no versions")
        vid = vers[0]["id"]
        r = requests.post(f"{API}/sites/{SITE_A}/menu/versions/{vid}/restore",
                          params={"vendor_id": VENDOR_1},
                          headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 403


# ---------- Regression: existing replace/append/clear still works ----------
class TestRegressionReplaceMode:
    def test_replace_mode_via_admin_upload(self, tokens):
        # Just verify endpoint exists and returns 200
        rows = [{"name": "REGR_A", "category": "Snacks", "price": 10, "is_vegetarian": True},
                {"name": "REGR_B", "category": "Snacks", "price": 12, "is_vegetarian": True}]
        xlsx = _make_xlsx(rows)
        files = {"file": ("r.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(
            f"{API}/sites/{SITE_A}/menu/upload-excel",
            params={"vendor_id": VENDOR_1, "mode": "replace"},
            headers=_h(tokens["master"]), files=files, timeout=30,
        )
        assert r.status_code == 200, r.text
        # Upload same file again → no duplicates
        files2 = {"file": ("r.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r2 = requests.post(
            f"{API}/sites/{SITE_A}/menu/upload-excel",
            params={"vendor_id": VENDOR_1, "mode": "replace"},
            headers=_h(tokens["master"]), files=files2, timeout=30,
        )
        assert r2.status_code == 200
        live = _live_menu(tokens["master"])
        names = [i["name"] for i in live]
        assert names.count("REGR_A") == 1
        assert names.count("REGR_B") == 1
