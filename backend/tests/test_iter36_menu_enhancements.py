"""Iter 36: Backend tests for 4 menu enhancements on top of Iter-35 bulk approval:
   1) Counter column in Excel template + parsing (vendor submit + admin direct upload)
   2) Per-Row Edit Before Approve (PATCH /admin/menu-uploads/{id}/items)
   3) Approval History Log (GET /admin/menu-uploads?status=decided)
   4) Approval Notifications (best-effort — 200 regardless)
"""
import io
import os
import pytest
import requests
from openpyxl import Workbook, load_workbook


def _load_env_url():
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if v:
        return v.rstrip("/")
    with open("/app/frontend/.env") as f:
        for ln in f:
            if ln.startswith("REACT_APP_BACKEND_URL="):
                return ln.strip().split("=", 1)[1].rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL not found")


BASE_URL = _load_env_url()
API = f"{BASE_URL}/api"

SITE_A = "6a96d37822127a5e04873e36"
VENDOR_1 = "6a96d37822127a5e04873e3c"

CREDS = {
    "master": ("admin@cravitoo.com", "admin123"),
    "vendor1": ("audit_vendor1@corpa.com", "Audit#1234"),
    "emp_a": ("audit_emp_a@corpa.com", "Audit#1234"),
}


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed {email}: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def tokens():
    return {k: _login(*v) for k, v in CREDS.items()}


def _h(t):
    return {"Authorization": f"Bearer {t}"}


def _live_menu(tok):
    r = requests.get(f"{API}/menu/{VENDOR_1}", headers=_h(tok), timeout=30)
    assert r.status_code == 200, r.text
    return r.json()


def _make_xlsx_with_counter(rows, include_counter=True):
    wb = Workbook()
    ws = wb.active
    headers = ["name", "description", "category", "price", "is_vegetarian",
               "image_url", "meal_periods"]
    if include_counter:
        headers.append("counter")
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------- Feature 1: Counter column ----------
class TestCounterColumn:
    def test_template_has_counter_column(self, tokens):
        r = requests.get(f"{API}/admin/menu-excel-template",
                         headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert "counter" in headers, f"counter col missing; got {headers}"
        assert len(headers) == 8, f"expected 8 headers incl counter, got {len(headers)}: {headers}"

    def test_vendor_submit_with_counter_then_admin_approve(self, tokens):
        rows = [
            {"name": "ITER36_CtrDosa", "description": "d", "category": "Breakfast",
             "price": 55, "is_vegetarian": True, "counter": "Counter 2"},
            {"name": "ITER36_CtrIdli", "description": "d", "category": "Breakfast",
             "price": 35, "is_vegetarian": True, "counter": "Counter 1"},
        ]
        xlsx = _make_xlsx_with_counter(rows)
        files = {"file": ("ctr.xlsx", xlsx,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/vendor/menu-uploads",
                          params={"site_id": SITE_A},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        assert r.status_code == 200, r.text
        req_id = r.json()["id"]

        # Approve
        rr = requests.post(f"{API}/admin/menu-uploads/{req_id}/approve",
                           headers=_h(tokens["master"]), timeout=30)
        assert rr.status_code == 200, rr.text

        # Verify live menu carries counter values
        live = _live_menu(tokens["master"])
        by_name = {i["name"]: i for i in live}
        assert "ITER36_CtrDosa" in by_name, f"missing item; live names={list(by_name)}"
        assert "counter" in by_name["ITER36_CtrDosa"], f"counter field missing in menu response: {by_name['ITER36_CtrDosa']}"
        assert by_name["ITER36_CtrDosa"]["counter"] == "Counter 2"
        assert by_name["ITER36_CtrIdli"]["counter"] == "Counter 1"

    def test_admin_direct_upload_with_counter(self, tokens):
        rows = [
            {"name": "ITER36_DirectCtr", "category": "Snacks", "price": 25,
             "is_vegetarian": True, "counter": "Counter 3"},
        ]
        xlsx = _make_xlsx_with_counter(rows)
        files = {"file": ("d.xlsx", xlsx,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/sites/{SITE_A}/menu/upload-excel",
                          params={"vendor_id": VENDOR_1, "mode": "append"},
                          headers=_h(tokens["master"]), files=files, timeout=30)
        assert r.status_code == 200, r.text
        live = _live_menu(tokens["master"])
        by_name = {i["name"]: i for i in live}
        assert "ITER36_DirectCtr" in by_name
        assert by_name["ITER36_DirectCtr"].get("counter") == "Counter 3"


# ---------- Feature 2: Per-Row Edit Before Approve ----------
class TestPerRowEdit:
    @pytest.fixture(scope="class")
    def pending_id(self, tokens):
        rows = [
            {"name": "EDIT_A", "category": "Snacks", "price": 100, "is_vegetarian": True, "counter": "Counter 1"},
            {"name": "EDIT_B", "category": "Snacks", "price": 200, "is_vegetarian": True, "counter": "Counter 1"},
            {"name": "EDIT_C", "category": "Snacks", "price": 300, "is_vegetarian": True, "counter": "Counter 1"},
        ]
        xlsx = _make_xlsx_with_counter(rows)
        files = {"file": ("e.xlsx", xlsx,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/vendor/menu-uploads",
                          params={"site_id": SITE_A},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        assert r.status_code == 200, r.text
        return r.json()["id"]

    def test_vendor_cannot_edit(self, tokens, pending_id):
        r = requests.patch(f"{API}/admin/menu-uploads/{pending_id}/items",
                           headers=_h(tokens["vendor1"]),
                           json={"items": [{"name": "EDIT_A", "price": 1}]}, timeout=30)
        assert r.status_code == 403

    def test_employee_cannot_edit(self, tokens, pending_id):
        r = requests.patch(f"{API}/admin/menu-uploads/{pending_id}/items",
                           headers=_h(tokens["emp_a"]),
                           json={"items": [{"name": "EDIT_A", "price": 1}]}, timeout=30)
        assert r.status_code == 403

    def test_master_edits_then_approves(self, tokens, pending_id):
        # Drop EDIT_C, change EDIT_A price to 80, keep EDIT_B unchanged
        payload = {"items": [
            {"name": "EDIT_A", "category": "Snacks", "price": 80,
             "is_vegetarian": True, "counter": "Counter 5"},
            {"name": "EDIT_B", "category": "Snacks", "price": 200,
             "is_vegetarian": True, "counter": "Counter 1"},
        ]}
        r = requests.patch(f"{API}/admin/menu-uploads/{pending_id}/items",
                           headers=_h(tokens["master"]), json=payload, timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["item_count"] == 2
        names = [i["name"] for i in data["items"]]
        assert set(names) == {"EDIT_A", "EDIT_B"}

        # Approve
        rr = requests.post(f"{API}/admin/menu-uploads/{pending_id}/approve",
                           headers=_h(tokens["master"]), timeout=30)
        assert rr.status_code == 200, rr.text
        assert rr.json()["applied"] == 2

        # Verify live menu reflects edits + dropped item
        live = _live_menu(tokens["master"])
        by_name = {i["name"]: i for i in live}
        assert "EDIT_A" in by_name
        assert by_name["EDIT_A"]["price"] == 80
        assert by_name["EDIT_A"].get("counter") == "Counter 5"
        assert "EDIT_B" in by_name
        assert "EDIT_C" not in by_name, "dropped item should not be in live menu"

    def test_cannot_edit_after_approve(self, tokens, pending_id):
        r = requests.patch(f"{API}/admin/menu-uploads/{pending_id}/items",
                           headers=_h(tokens["master"]),
                           json={"items": [{"name": "X", "price": 1}]}, timeout=30)
        assert r.status_code == 400  # already approved


# ---------- Feature 3: Approval History Log ----------
class TestApprovalHistory:
    def test_history_returns_decided(self, tokens):
        r = requests.get(f"{API}/admin/menu-uploads",
                         params={"status": "decided", "site_id": SITE_A},
                         headers=_h(tokens["master"]), timeout=30)
        assert r.status_code == 200, r.text
        rows = r.json()
        assert isinstance(rows, list) and len(rows) >= 1
        # All rows must be approved or rejected
        for row in rows:
            assert row["status"] in ("approved", "rejected"), row
            for k in ("decided_by", "decided_at", "vendor_name", "item_count"):
                assert k in row, f"missing {k}: {row}"
            assert row["decided_by"], f"decided_by empty: {row}"

    def test_history_vendor_forbidden(self, tokens):
        r = requests.get(f"{API}/admin/menu-uploads",
                         params={"status": "decided"},
                         headers=_h(tokens["vendor1"]), timeout=30)
        assert r.status_code == 403


# ---------- Feature 4: Notifications (best-effort — should not break API) ----------
class TestNotificationsBestEffort:
    def test_approve_returns_200_even_with_email(self, tokens):
        rows = [{"name": "NOTIFY_Approve", "category": "Snacks", "price": 10, "is_vegetarian": True}]
        xlsx = _make_xlsx_with_counter(rows)
        files = {"file": ("n.xlsx", xlsx,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/vendor/menu-uploads",
                          params={"site_id": SITE_A},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        assert r.status_code == 200
        rid = r.json()["id"]
        rr = requests.post(f"{API}/admin/menu-uploads/{rid}/approve",
                           headers=_h(tokens["master"]), timeout=30)
        assert rr.status_code == 200, rr.text
        assert rr.json()["status"] == "approved"

    def test_reject_returns_200_with_note(self, tokens):
        rows = [{"name": "NOTIFY_Reject", "category": "Snacks", "price": 10, "is_vegetarian": True}]
        xlsx = _make_xlsx_with_counter(rows)
        files = {"file": ("nr.xlsx", xlsx,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/vendor/menu-uploads",
                          params={"site_id": SITE_A},
                          headers=_h(tokens["vendor1"]), files=files, timeout=30)
        assert r.status_code == 200
        rid = r.json()["id"]
        rr = requests.post(f"{API}/admin/menu-uploads/{rid}/reject",
                           headers=_h(tokens["master"]),
                           json={"note": "Please reduce prices"}, timeout=30)
        assert rr.status_code == 200, rr.text
        assert rr.json()["status"] == "rejected"
        # Verify note stored via decided history
        hist = requests.get(f"{API}/admin/menu-uploads",
                            params={"status": "decided", "site_id": SITE_A},
                            headers=_h(tokens["master"]), timeout=30).json()
        row = next((x for x in hist if x["id"] == rid), None)
        assert row is not None
        assert row["decision_note"] == "Please reduce prices"


# ---------- Regression: full flow no duplicates + RBAC ----------
class TestRegression:
    def test_no_duplicates_on_repeat_approve_cycle(self, tokens):
        # Submit + approve same file twice via new submissions; replace mode should NOT duplicate
        rows = [{"name": "REG36_Alpha", "category": "Snacks", "price": 15, "is_vegetarian": True}]
        xlsx = _make_xlsx_with_counter(rows)
        for _ in range(2):
            files = {"file": ("r.xlsx", xlsx,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
            r = requests.post(f"{API}/vendor/menu-uploads",
                              params={"site_id": SITE_A},
                              headers=_h(tokens["vendor1"]), files=files, timeout=30)
            rid = r.json()["id"]
            requests.post(f"{API}/admin/menu-uploads/{rid}/approve",
                          headers=_h(tokens["master"]), timeout=30)
        live = _live_menu(tokens["master"])
        names = [i["name"] for i in live]
        assert names.count("REG36_Alpha") == 1, f"duplicate found: {names}"

    def test_vendor_cannot_use_admin_endpoints(self, tokens):
        # list history
        assert requests.get(f"{API}/admin/menu-uploads",
                            headers=_h(tokens["vendor1"]), timeout=30).status_code == 403
