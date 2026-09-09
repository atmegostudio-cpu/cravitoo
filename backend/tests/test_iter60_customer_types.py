"""Iter60: Customer Types admin + Non-Corporate Filter regression."""
import os
import pytest
import requests
from dotenv import load_dotenv

load_dotenv("/app/frontend/.env")
BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")

MASTER = {"email": "admin@cravitoo.com", "password": "admin123"}


@pytest.fixture(scope="module")
def master_client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    r = s.post(f"{BASE_URL}/api/auth/login", json=MASTER, timeout=30)
    assert r.status_code == 200, r.text
    tok = r.json().get("access_token") or r.json().get("token")
    if tok:
        s.headers["Authorization"] = f"Bearer {tok}"
    return s


# ─── Customer-types GET ───────────────────────────────────────────────────
def test_list_customer_types_authenticated(master_client):
    r = master_client.get(f"{BASE_URL}/api/customer-types", timeout=15)
    assert r.status_code == 200
    data = r.json()
    assert isinstance(data, list) and len(data) >= 5
    names = {d["name"] for d in data}
    assert {"Guest", "Housekeeping", "Security", "Drivers", "Facility Management"} <= names
    for d in data:
        assert "id" in d and "name" in d


def test_list_customer_types_unauthenticated():
    r = requests.get(f"{BASE_URL}/api/customer-types", timeout=15)
    assert r.status_code in (401, 403)


# ─── Admin CRUD as master_admin ───────────────────────────────────────────
_new_id = {"id": None}


def test_add_customer_type(master_client):
    r = master_client.post(f"{BASE_URL}/api/admin/customer-types",
                           json={"name": "TEST_Interns60"}, timeout=15)
    assert r.status_code == 200, r.text
    lst = master_client.get(f"{BASE_URL}/api/customer-types", timeout=15).json()
    match = [d for d in lst if d["name"] == "TEST_Interns60"]
    assert match, "new type not persisted"
    _new_id["id"] = match[0]["id"]


def test_rename_customer_type(master_client):
    assert _new_id["id"]
    r = master_client.patch(f"{BASE_URL}/api/admin/customer-types/{_new_id['id']}",
                            json={"name": "TEST_Interns60_Renamed"}, timeout=15)
    assert r.status_code == 200, r.text
    lst = master_client.get(f"{BASE_URL}/api/customer-types", timeout=15).json()
    names = {d["name"] for d in lst}
    assert "TEST_Interns60_Renamed" in names
    assert "TEST_Interns60" not in names


def test_delete_customer_type(master_client):
    assert _new_id["id"]
    r = master_client.delete(f"{BASE_URL}/api/admin/customer-types/{_new_id['id']}", timeout=15)
    assert r.status_code == 200
    lst = master_client.get(f"{BASE_URL}/api/customer-types", timeout=15).json()
    assert "TEST_Interns60_Renamed" not in {d["name"] for d in lst}


# ─── RBAC: non-master forbidden ───────────────────────────────────────────
@pytest.fixture(scope="module")
def non_master_client():
    """Login as a non-master role. Use the audit corporate admin."""
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    r = s.post(f"{BASE_URL}/api/auth/login",
               json={"email": "test_nonmaster@x.com", "password": "t123"},
               timeout=30)
    if r.status_code != 200:
        pytest.skip(f"non-master seed missing: {r.status_code} {r.text[:120]}")
    tok = r.json().get("access_token") or r.json().get("token")
    if tok:
        s.headers["Authorization"] = f"Bearer {tok}"
    return s


def test_non_master_get_types_allowed(non_master_client):
    r = non_master_client.get(f"{BASE_URL}/api/customer-types", timeout=15)
    assert r.status_code == 200


def test_non_master_add_forbidden(non_master_client):
    r = non_master_client.post(f"{BASE_URL}/api/admin/customer-types",
                               json={"name": "TEST_Blocked"}, timeout=15)
    assert r.status_code == 403


def test_non_master_patch_forbidden(non_master_client):
    r = non_master_client.patch(f"{BASE_URL}/api/admin/customer-types/000000000000000000000000",
                                json={"name": "X"}, timeout=15)
    assert r.status_code == 403


def test_non_master_delete_forbidden(non_master_client):
    r = non_master_client.delete(
        f"{BASE_URL}/api/admin/customer-types/000000000000000000000000", timeout=15)
    assert r.status_code == 403


# ─── Sales report filters include customer_types ─────────────────────────
def test_sales_filters_have_customer_types(master_client):
    r = master_client.get(f"{BASE_URL}/api/admin/sales-report/filters", timeout=20)
    assert r.status_code == 200
    body = r.json()
    assert "customer_types" in body
    names = {c["name"] for c in body["customer_types"]}
    assert "Corporate (employees)" in names
    # active types should also be included
    seeded = {"Guest", "Housekeeping", "Security", "Drivers", "Facility Management"}
    assert seeded <= names


# ─── Sales report customer_type_summary + filter param ────────────────────
def test_sales_report_customer_type_summary(master_client):
    r = master_client.get(f"{BASE_URL}/api/admin/sales-report?month=2026-06", timeout=30)
    assert r.status_code == 200
    body = r.json()
    assert "customer_type_summary" in body
    assert isinstance(body["customer_type_summary"], list)
    # legacy orders should bucket to Corporate
    labels = {row["customer_type"] for row in body["customer_type_summary"]}
    # Not asserting Corporate strictly (depends on data), but structure check:
    for row in body["customer_type_summary"]:
        assert "customer_type" in row and "total" in row
        assert isinstance(row["total"], (int, float))


def test_sales_report_customer_types_filter(master_client):
    r = master_client.get(
        f"{BASE_URL}/api/admin/sales-report?month=2026-06&customer_types=Corporate", timeout=30)
    assert r.status_code == 200
    body = r.json()
    for row in body.get("customer_type_summary", []):
        assert row["customer_type"] == "Corporate"
    for o in body.get("orders", []):
        assert o.get("customer_type") == "Corporate"


def test_sales_report_orders_have_customer_type_field(master_client):
    r = master_client.get(f"{BASE_URL}/api/admin/sales-report?month=2026-06", timeout=30)
    body = r.json()
    for o in body.get("orders", [])[:5]:
        assert "customer_type" in o


# ─── Excel export contains Customer Types sheet ──────────────────────────
def test_excel_export_has_customer_types_sheet(master_client, tmp_path):
    r = master_client.get(
        f"{BASE_URL}/api/admin/sales-report?month=2026-06&format=xlsx", timeout=45)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    p = tmp_path / "sales.xlsx"
    p.write_bytes(r.content)
    import openpyxl
    wb = openpyxl.load_workbook(str(p))
    assert "Customer Types" in wb.sheetnames
    ws = wb["Orders"]
    headers = [c.value for c in ws[1]]
    assert "Customer Type" in headers


# ─── Vendor orders projection carries customer_type + is_manual ──────────
def test_orders_projection_has_customer_type_and_is_manual(master_client):
    r = master_client.get(f"{BASE_URL}/api/orders", timeout=30)
    if r.status_code != 200:
        pytest.skip(f"/api/orders not accessible for master: {r.status_code}")
    orders = r.json()
    if not orders:
        pytest.skip("no orders in DB")
    o = orders[0]
    # Fields may be absent for legacy docs, but the projection must at least allow them
    # (i.e., presence for a manual order). We can't guarantee a manual order exists, so
    # just check that if any manual order exists, customer_type is present.
    manual = [x for x in orders if x.get("is_manual")]
    for m in manual:
        assert "customer_type" in m and m["customer_type"]
