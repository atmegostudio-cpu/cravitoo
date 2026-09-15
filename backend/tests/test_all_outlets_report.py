"""Backend tests for Total Sales (All Outlets) report + Sales Report per-site timezone.

Covers:
  - /api/vendor/all-outlets-report (month + range)
  - /api/vendor/all-outlets-report/export (xlsx, 3 sheets)
  - 403 for non-operator vendor and non-vendor role
  - /api/vendor/reports/sales-orders returns site_timezone per row
  - /api/vendor/reports/sales-export?format=csv header includes 'site local'
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

def _load_backend_url():
    url = os.environ.get("REACT_APP_BACKEND_URL")
    if not url:
        try:
            with open("/app/frontend/.env") as fh:
                for line in fh:
                    if line.startswith("REACT_APP_BACKEND_URL="):
                        url = line.split("=", 1)[1].strip()
                        break
        except FileNotFoundError:
            pass
    if not url:
        raise RuntimeError("REACT_APP_BACKEND_URL not set")
    return url.rstrip("/")

BASE_URL = _load_backend_url()

OPERATOR = ("tz_operator@cravitoo.com", "Pass1234")
ADMIN = ("admin@cravitoo.com", "admin123")


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed for {email}: {r.status_code} {r.text}"
    return s


@pytest.fixture(scope="module")
def operator_session():
    return _login(*OPERATOR)


@pytest.fixture(scope="module")
def admin_session():
    return _login(*ADMIN)


# ---------- 1) all-outlets-report month + range ----------
class TestAllOutletsReport:
    def test_report_month_2026_09(self, operator_session):
        r = operator_session.get(f"{BASE_URL}/api/vendor/all-outlets-report", params={"month": "2026-09"}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert "summary" in d and "per_vendor" in d and "per_counter" in d and "outlets" in d
        s = d["summary"]
        # Main agent reported: orders=2, total=330, paid=210, pending=120, avg=165 for Sept 2026
        assert s["total_orders"] == 2
        assert s["total_amount"] == 330
        assert s["paid_amount"] == 210
        assert s["pending_amount"] == 120
        assert s["avg_order_value"] == 165
        assert len(d["per_vendor"]) >= 1
        assert len(d["per_counter"]) >= 1

    def test_report_range_2026(self, operator_session):
        r = operator_session.get(
            f"{BASE_URL}/api/vendor/all-outlets-report",
            params={"from": "2026-01-01T00:00:00Z", "to": "2026-12-31T23:59:59Z"},
            timeout=30,
        )
        assert r.status_code == 200
        d = r.json()
        assert d["summary"]["total_orders"] >= 2  # wider range >= narrow month
        assert d["summary"]["total_amount"] >= 330

    def test_range_larger_than_month(self, operator_session):
        m = operator_session.get(f"{BASE_URL}/api/vendor/all-outlets-report", params={"month": "2026-09"}).json()
        r = operator_session.get(
            f"{BASE_URL}/api/vendor/all-outlets-report",
            params={"from": "2026-01-01T00:00:00Z", "to": "2026-12-31T23:59:59Z"},
        ).json()
        assert r["summary"]["total_orders"] >= m["summary"]["total_orders"]


# ---------- 2) export xlsx ----------
class TestAllOutletsExport:
    def test_export_returns_xlsx_with_3_sheets(self, operator_session):
        r = operator_session.get(
            f"{BASE_URL}/api/vendor/all-outlets-report/export",
            params={"month": "2026-09"}, timeout=30,
        )
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Summary", "By Vendor", "By Counter"], wb.sheetnames
        # Summary values must match the JSON report
        summary = operator_session.get(
            f"{BASE_URL}/api/vendor/all-outlets-report", params={"month": "2026-09"}
        ).json()["summary"]
        ws = wb["Summary"]
        vals = {row[0].value: row[1].value for row in ws.iter_rows(min_row=1, max_row=ws.max_row) if row[0].value}
        assert vals.get("Total Orders") == summary["total_orders"]
        assert float(vals.get("Total Sales (INR)")) == float(summary["total_amount"])
        assert float(vals.get("Paid (INR)")) == float(summary["paid_amount"])


# ---------- 3) 403 permission checks ----------
class TestPermission403:
    def test_non_vendor_role_gets_403(self, admin_session):
        r = admin_session.get(f"{BASE_URL}/api/vendor/all-outlets-report", params={"month": "2026-09"})
        assert r.status_code == 403
        r2 = admin_session.get(f"{BASE_URL}/api/vendor/all-outlets-report/export", params={"month": "2026-09"})
        assert r2.status_code == 403

    def test_plain_vendor_no_assigned_vendors_gets_403(self):
        s = _login("audit_vendor1@corpa.com", "Audit#1234")
        r = s.get(f"{BASE_URL}/api/vendor/all-outlets-report", params={"month": "2026-09"})
        assert r.status_code == 403
        r2 = s.get(f"{BASE_URL}/api/vendor/all-outlets-report/export", params={"month": "2026-09"})
        assert r2.status_code == 403

    def test_unauthenticated_401_or_403(self):
        r = requests.get(f"{BASE_URL}/api/vendor/all-outlets-report", params={"month": "2026-09"})
        assert r.status_code in (401, 403)


# ---------- 4) sales-orders returns site_timezone ----------
class TestSalesOrdersTimezone:
    def test_sales_orders_include_site_timezone_field(self, operator_session):
        r = operator_session.get(
            f"{BASE_URL}/api/vendor/reports/sales-orders",
            params={"from": "2026-01-01T00:00:00Z", "to": "2026-12-31T23:59:59Z", "size": 50},
            timeout=30,
        )
        assert r.status_code == 200
        rows = r.json().get("rows", [])
        assert rows, "expected at least 1 row"
        for row in rows:
            assert "site_timezone" in row
            assert isinstance(row["site_timezone"], str) and row["site_timezone"]

    def test_csv_export_header_and_site_local(self, operator_session):
        r = operator_session.get(
            f"{BASE_URL}/api/vendor/reports/sales-export",
            params={"from": "2026-01-01T00:00:00Z", "to": "2026-12-31T23:59:59Z", "format": "csv"},
            timeout=30,
        )
        assert r.status_code == 200
        text = r.text
        first_line = text.splitlines()[0]
        assert "Date & Time (site local)" in first_line, first_line
