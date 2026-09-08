"""Admin Sales Report — role-scoped totals split by Client, City, Site and
Vendor, with cascading multi-select filters (Client -> City -> Site -> Vendor)
plus Date / Date-range / Month filters and Excel (.xlsx) export.

Access: master_admin (all sites), corporate_admin (their company's sites),
site_admin (their site only).
"""
import io
from datetime import datetime, timezone, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse


def _parse_range(date: Optional[str], start: Optional[str], end: Optional[str], month: Optional[str]):
    """Return (start_dt, end_dt) UTC-aware, inclusive start / exclusive end."""
    def d(s):
        return datetime.fromisoformat(s).replace(tzinfo=timezone.utc)
    if month:  # 'YYYY-MM'
        y, m = int(month[:4]), int(month[5:7])
        s = datetime(y, m, 1, tzinfo=timezone.utc)
        e = datetime(y + (m // 12), (m % 12) + 1, 1, tzinfo=timezone.utc)
        return s, e
    if date:
        s = d(date)
        return s, s + timedelta(days=1)
    if start and end:
        return d(start), d(end) + timedelta(days=1)
    # default: last 30 days
    now = datetime.now(timezone.utc)
    return now - timedelta(days=30), now + timedelta(days=1)


def _split_ids(v: Optional[str]):
    """Comma-separated id string -> clean list (empty -> [])."""
    if not v:
        return []
    return [p.strip() for p in v.split(",") if p.strip()]


def make_router(db, safe_objectid, get_current_user):
    r = APIRouter()

    async def _allowed_site_ids(user):
        role = user.get("role")
        if role == "master_admin":
            return None  # all
        if role == "site_admin":
            return [user.get("site_id")] if user.get("site_id") else []
        if role == "corporate_admin":
            cid = user.get("company_id")
            sites = await db.sites.find({"company_id": cid}, {"_id": 1}).to_list(1000)
            return [str(s["_id"]) for s in sites]
        raise HTTPException(status_code=403, detail="Not allowed")

    async def _candidate_sites(user):
        """Site metadata within the caller's role scope."""
        allowed = await _allowed_site_ids(user)
        site_query = {}
        if allowed is not None:
            oids = [safe_objectid(x, "Site") for x in allowed if x]
            if not oids:
                return allowed, {}
            site_query["_id"] = {"$in": oids}
        docs = await db.sites.find(site_query, {"name": 1, "city": 1, "city_id": 1, "company_id": 1}).to_list(5000)
        meta = {}
        for sd in docs:
            sid = str(sd["_id"])
            meta[sid] = {
                "name": sd.get("name") or sid,
                "city": sd.get("city"),
                "city_id": str(sd["city_id"]) if sd.get("city_id") else None,
                "company_id": str(sd["company_id"]) if sd.get("company_id") else None,
            }
        return allowed, meta

    @r.get("/admin/sales-report/filters")
    async def sales_report_filters(user: dict = Depends(get_current_user)):
        """Cascading filter options scoped to the caller's role."""
        allowed, meta = await _candidate_sites(user)

        comp_ids = {m["company_id"] for m in meta.values() if m["company_id"]}
        city_ids = {m["city_id"] for m in meta.values() if m["city_id"]}

        clients = []
        for cid in comp_ids:
            co = await db.companies.find_one({"_id": safe_objectid(cid, "Company")}, {"name": 1})
            if co:
                clients.append({"id": cid, "name": co.get("name") or cid})
        clients.sort(key=lambda x: (x["name"] or "").lower())

        cities = []
        for cid in city_ids:
            ci = await db.cities.find_one({"_id": safe_objectid(cid, "City")}, {"name": 1})
            cities.append({"id": cid, "name": (ci or {}).get("name") or cid})
        cities.sort(key=lambda x: (x["name"] or "").lower())

        sites = [{"id": sid, "name": m["name"], "company_id": m["company_id"],
                  "city_id": m["city_id"], "city": m["city"]}
                 for sid, m in meta.items()]
        sites.sort(key=lambda x: (x["name"] or "").lower())

        # vendors mapped to the candidate sites (+ which sites each serves)
        site_id_strs = list(meta.keys())
        oid_list = [safe_objectid(x, "Site") for x in site_id_strs]
        maps = await db.vendor_site_mappings.find(
            {"site_id": {"$in": site_id_strs + oid_list}}).to_list(20000)
        vendor_sites = {}
        for mp in maps:
            vid = str(mp.get("vendor_id")) if mp.get("vendor_id") else None
            sid = str(mp.get("site_id")) if mp.get("site_id") else None
            if vid and sid:
                vendor_sites.setdefault(vid, set()).add(sid)

        vendors = []
        for vid, sids in vendor_sites.items():
            vdoc = await db.vendors.find_one({"_id": safe_objectid(vid, "Vendor")}, {"name": 1})
            if vdoc:
                vendors.append({"id": vid, "name": vdoc.get("name") or vid, "site_ids": sorted(sids)})
        vendors.sort(key=lambda x: (x["name"] or "").lower())

        return {"clients": clients, "cities": cities, "sites": sites, "vendors": vendors}

    async def _gather(user, date, start, end, month, client_ids, city_ids, site_ids, vendor_ids):
        allowed, meta = await _candidate_sites(user)
        s, e = _parse_range(date, start, end, month)

        # resolve name maps for clients + cities in scope
        comp_ids = {m["company_id"] for m in meta.values() if m["company_id"]}
        city_ids_scope = {m["city_id"] for m in meta.values() if m["city_id"]}
        company_names, city_names = {}, {}
        for cid in comp_ids:
            co = await db.companies.find_one({"_id": safe_objectid(cid, "Company")}, {"name": 1})
            if co:
                company_names[cid] = co.get("name")
        for cid in city_ids_scope:
            ci = await db.cities.find_one({"_id": safe_objectid(cid, "City")}, {"name": 1})
            if ci:
                city_names[cid] = ci.get("name")

        def site_city_label(sid):
            m = meta.get(sid) or {}
            return m.get("city") or city_names.get(m.get("city_id")) or "—"

        def site_client_label(sid):
            m = meta.get(sid) or {}
            return company_names.get(m.get("company_id")) or "—"

        # effective site set after Client / City / Site filters
        effective = set(meta.keys())
        if client_ids:
            effective = {sid for sid in effective if (meta[sid].get("company_id") in client_ids)}
        if city_ids:
            effective = {sid for sid in effective if (meta[sid].get("city_id") in city_ids)}
        if site_ids:
            effective &= set(site_ids)

        apply_site_filter = (allowed is not None) or bool(client_ids or city_ids or site_ids)
        q = {"created_at": {"$gte": s, "$lt": e}, "status": {"$ne": "cancelled"}}
        if apply_site_filter:
            q["site_id"] = {"$in": list(effective)}
        if vendor_ids:
            q["vendor_id"] = {"$in": vendor_ids}

        orders = await db.orders.find(q).sort("created_at", -1).to_list(100000)

        # vendor name lookups (orders may reference vendors outside meta scope)
        vend_ids = {o.get("vendor_id") for o in orders if o.get("vendor_id")}
        vend_names = {}
        for vid in vend_ids:
            doc = await db.vendors.find_one({"_id": safe_objectid(vid, "Vendor")}, {"name": 1})
            if doc:
                vend_names[vid] = doc.get("name")

        rows = []
        site_tot, city_tot, vend_tot, client_tot, grand = {}, {}, {}, {}, 0.0
        for o in orders:
            amt = float(o.get("total_amount") or 0)
            grand += amt
            sid, vid = o.get("site_id"), o.get("vendor_id")
            site_name = (meta.get(sid) or {}).get("name") or (sid or "—")
            city_label = site_city_label(sid)
            client_label = site_client_label(sid)
            site_tot[sid] = site_tot.get(sid, 0.0) + amt
            city_tot[city_label] = city_tot.get(city_label, 0.0) + amt
            client_tot[client_label] = client_tot.get(client_label, 0.0) + amt
            vend_tot[(sid, vid)] = vend_tot.get((sid, vid), 0.0) + amt
            ca = o.get("created_at")
            rows.append({
                "order_id": str(o["_id"]),
                "date": ca.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M") if ca else "",
                "client": client_label,
                "city": city_label,
                "site": site_name,
                "vendor": vend_names.get(vid, vid or "—"),
                "employee": o.get("employee_name") or o.get("employee_email") or "",
                "items": ", ".join(f"{i.get('quantity', 1)}× {i.get('name', '')}" for i in (o.get("items") or [])),
                "amount": round(amt, 2),
                "payment_status": o.get("payment_status") or "",
                "payment_method": o.get("payment_method") or o.get("payment_mode") or "",
                "fulfilment": o.get("status") or "",
            })

        site_summary = [{"site": (meta.get(k) or {}).get("name") or (k or "—"),
                         "city": site_city_label(k), "client": site_client_label(k),
                         "total": round(v, 2)}
                        for k, v in sorted(site_tot.items(), key=lambda x: -x[1])]
        city_summary = [{"city": k, "total": round(v, 2)}
                        for k, v in sorted(city_tot.items(), key=lambda x: -x[1])]
        client_summary = [{"client": k, "total": round(v, 2)}
                          for k, v in sorted(client_tot.items(), key=lambda x: -x[1])]
        vendor_summary = [{"site": (meta.get(k[0]) or {}).get("name") or (k[0] or "—"),
                           "city": site_city_label(k[0]),
                           "vendor": vend_names.get(k[1], k[1] or "—"),
                           "total": round(v, 2)}
                          for k, v in sorted(vend_tot.items(), key=lambda x: -x[1])]

        return {
            "range": {"start": s.isoformat(), "end": e.isoformat()},
            "grand_total": round(grand, 2),
            "order_count": len(rows),
            "client_summary": client_summary,
            "city_summary": city_summary,
            "site_summary": site_summary,
            "vendor_summary": vendor_summary,
            "orders": rows,
        }

    @r.get("/admin/sales-report")
    async def sales_report(
        date: Optional[str] = None,
        start: Optional[str] = None,
        end: Optional[str] = None,
        month: Optional[str] = None,
        client_ids: Optional[str] = None,
        city_ids: Optional[str] = None,
        site_ids: Optional[str] = None,
        vendor_ids: Optional[str] = None,
        format: str = Query("json"),
        user: dict = Depends(get_current_user),
    ):
        data = await _gather(
            user, date, start, end, month,
            _split_ids(client_ids), _split_ids(city_ids),
            _split_ids(site_ids), _split_ids(vendor_ids),
        )
        if format != "xlsx":
            return data

        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        # Sheet 1: Orders
        ws = wb.active
        ws.title = "Orders"
        headers = ["Date", "Client", "City", "Site", "Vendor", "Employee", "Items", "Amount",
                   "Payment Status", "Payment Method", "Fulfilment", "Order ID"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for o in data["orders"]:
            ws.append([o["date"], o["client"], o["city"], o["site"], o["vendor"], o["employee"],
                       o["items"], o["amount"], o["payment_status"], o["payment_method"],
                       o["fulfilment"], o["order_id"]])
        # Sheet 2: City Totals
        wsc = wb.create_sheet("City Totals")
        wsc.append(["City", "Total (₹)"])
        for c in wsc[1]:
            c.font = Font(bold=True)
        for row in data["city_summary"]:
            wsc.append([row["city"], row["total"]])
        wsc.append(["GRAND TOTAL", data["grand_total"]])
        wsc[wsc.max_row][0].font = Font(bold=True); wsc[wsc.max_row][1].font = Font(bold=True)
        # Sheet 3: Site Totals
        ws2 = wb.create_sheet("Site Totals")
        ws2.append(["Client", "City", "Site", "Total (₹)"])
        for c in ws2[1]:
            c.font = Font(bold=True)
        for row in data["site_summary"]:
            ws2.append([row.get("client", "—"), row.get("city", "—"), row["site"], row["total"]])
        ws2.append(["", "", "GRAND TOTAL", data["grand_total"]])
        for cell in ws2[ws2.max_row]:
            cell.font = Font(bold=True)
        # Sheet 4: Vendor Totals
        ws3 = wb.create_sheet("Vendor Totals")
        ws3.append(["City", "Site", "Vendor", "Total (₹)"])
        for c in ws3[1]:
            c.font = Font(bold=True)
        for v in data["vendor_summary"]:
            ws3.append([v.get("city", "—"), v["site"], v["vendor"], v["total"]])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"cravitoo-sales-{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{fname}"'},
        )

    return r
